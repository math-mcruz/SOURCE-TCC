#!/usr/bin/env python3
"""
extrair_pares_gabarito.py

Passo 2 do método (E4, seção 1.3): extração e filtragem dos pares
comentário-código da CBSA, para revisão conjunta discente-orientador.

Procedimento (heurístico, somente leitura na origem):
  1. Localiza a PROCEDURE DIVISION de cada programa (comentários de cabeçalho,
     anteriores a ela, não entram).
  2. Cada bloco contíguo de comentário (formato fixo: '*' ou '/' na coluna 7)
     é pareado ao código que o segue, até o próximo bloco de comentário com
     texto (limite de LIMITE_LINHAS linhas de código por trecho).
  3. Retém o par se o trecho contém estrutura de decisão (IF, EVALUATE, WHEN)
     ou tratamento de falha (ABEND, SQLCODE, HANDLE, DFHRESP, EIBRESP, EIBRCODE).
     Os demais vão para a aba "Descartados" (rastreabilidade do filtro).
  4. Gera a planilha de revisão (.xlsx) com colunas para revisão do discente,
     do orientador e a redação final da regra atômica.

Uso (Windows):
  pip install openpyxl
  python extrair_pares_gabarito.py --origem "...\\DATABASE\\ORIGINAL\\DATABASE\\cobol_src" --saida gabarito_candidatos.xlsx

Opcional: --programas DBCRFUN XFRFUN BNK1CRA BNK1TFN  (padrão: esses quatro)
"""

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

LIMITE_LINHAS = 30
PROGRAMAS_PADRAO = ["DBCRFUN", "XFRFUN", "BNK1CRA", "BNK1TFN"]

RE_PD = re.compile(r"^\s*PROCEDURE\s+DIVISION")
RE_LIT = re.compile(r"'[^']*'")
RE_IF = re.compile(r"(?<![-\w])IF(?![-\w])")
RE_EVAL = re.compile(r"(?<![-\w])EVALUATE(?![-\w])")
RE_WHEN = re.compile(r"(?<![-\w])WHEN(?![-\w])")
RE_FALHA = re.compile(r"(?<![-\w])(ABEND|SQLCODE|HANDLE|DFHRESP|EIBRESP|EIBRCODE)(?![-\w])")
RE_CONDICAO = re.compile(
    r"\b(if|when|unless|only|must|should|ensure|check|otherwise|cannot|not|sufficient|invalid|valid)\b",
    re.IGNORECASE,
)
RE_DESATIVADO = re.compile(
    r"^(MOVE|DISPLAY|PERFORM|EXEC|SET|COMPUTE|ADD|SUBTRACT|CALL|GO\s+TO|INITIALIZE|"
    r"END-\w+|IF|EVALUATE|SOURCE-COMPUTER|OBJECT-COMPUTER)\b"
)


def eh_comentario(l):
    return len(l) > 6 and l[6] in "*/"


def eh_debug(l):
    return len(l) > 6 and l[6] in "Dd"


def texto_comentario(l):
    """Texto de uma linha de comentário, sem indicador nem faixas de asteriscos."""
    return l[7:].strip().lstrip("*").strip()


def intervalo(a, b):
    return f"{a}" if a == b else f"{a}-{b}"


def classificar(linhas_codigo):
    """Retorna (tipo, estruturas) para o trecho; tipo None = procedural."""
    n_if = n_ev = n_wh = 0
    falha = False
    for l in linhas_codigo:
        if eh_debug(l):
            continue
        t = RE_LIT.sub("''", l[7:72])
        n_if += len(RE_IF.findall(t))
        n_ev += len(RE_EVAL.findall(t))
        n_wh += len(RE_WHEN.findall(t))
        falha = falha or bool(RE_FALHA.search(t))
    estruturas = []
    if n_if:
        estruturas.append("IF")
    if n_ev:
        estruturas.append("EVALUATE")
    if n_wh:
        estruturas.append("WHEN")
    if falha:
        estruturas.append("falha")
    decisao = (n_if + n_ev + n_wh) > 0
    if decisao and falha:
        tipo = "Decisão e falha"
    elif decisao:
        tipo = "Decisão"
    elif falha:
        tipo = "Falha"
    else:
        tipo = None
    return tipo, ", ".join(estruturas)


def extrair(caminho: Path):
    programa = caminho.stem.upper()
    dados = caminho.read_bytes().decode("latin-1").replace("\r", "")
    L = dados.split("\n")
    if L and L[-1] == "":
        L.pop()
    n = len(L)
    pd = next((i for i, l in enumerate(L) if not eh_comentario(l) and RE_PD.search(l)), None)
    if pd is None:
        return []

    pares = []
    seq = 0
    i = pd + 1
    while i < n:
        if not eh_comentario(L[i]):
            i += 1
            continue
        j = i
        while j < n and (eh_comentario(L[j]) or not L[j].strip()):
            j += 1
        com = [(k, texto_comentario(L[k])) for k in range(i, j) if eh_comentario(L[k])]
        com = [(k, t) for k, t in com if t]
        if not com:
            i = j
            continue

        fundidos = False
        cod = []
        while True:
            k = j
            while k < n:
                if eh_comentario(L[k]):
                    if texto_comentario(L[k]):
                        break
                    k += 1
                    continue
                if L[k].strip():
                    cod.append(k)
                k += 1
            ativo = any(not eh_debug(L[m]) for m in cod)
            if ativo or k >= n:
                break
            # sem código ativo (só depuração ou nada): o comentário seguinte
            # continua a mesma explicação; funde os blocos
            j2 = k
            while j2 < n and (eh_comentario(L[j2]) or not L[j2].strip()):
                j2 += 1
            extra = [(m, texto_comentario(L[m])) for m in range(k, j2) if eh_comentario(L[m])]
            com += [(m, t) for m, t in extra if t]
            fundidos = True
            j = j2
        truncado = len(cod) > LIMITE_LINHAS
        cod = cod[:LIMITE_LINHAS]

        seq += 1
        texto = " ".join(t for _, t in com)
        alertas = []
        if RE_DESATIVADO.match(com[0][1]):
            alertas.append("possível código desativado")
        if len(texto.split()) <= 3:
            alertas.append("comentário curto")
        if fundidos:
            alertas.append("comentários adjacentes fundidos")
        if truncado:
            alertas.append(f"trecho truncado ({LIMITE_LINHAS} linhas)")

        if cod:
            tipo, estruturas = classificar([L[m] for m in cod])
            codigo = "\n".join(
                ("[D] " if eh_debug(L[m]) else "") + L[m][7:72].rstrip() for m in cod
            )
            l_cod = intervalo(cod[0] + 1, cod[-1] + 1)
        else:
            tipo, estruturas, codigo, l_cod = None, "", "", ""

        pares.append(
            {
                "id": f"{programa}-{seq:03d}",
                "programa": programa,
                "l_com": intervalo(com[0][0] + 1, com[-1][0] + 1),
                "l_cod": l_cod,
                "comentario": texto,
                "codigo": codigo,
                "tipo": tipo,
                "estruturas": estruturas,
                "alertas": "; ".join(alertas) or None,
                "sem_codigo": not cod,
            }
        )
        i = k
    return pares


# ----------------------------------------------------------------- planilha
FONTE = "Arial"
F_NORMAL = Font(name=FONTE, size=10)
F_NEGRITO = Font(name=FONTE, size=10, bold=True)
F_CABEC = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
F_CODIGO = Font(name="Courier New", size=9)
F_TITULO = Font(name=FONTE, size=13, bold=True)
PREENCHE_CABEC = PatternFill("solid", fgColor="2F4F6F")
PREENCHE_ENTRADA = PatternFill("solid", fgColor="FFF9DB")
TOPO = Alignment(vertical="top", wrap_text=True)


def cabecalho(ws, titulos, larguras):
    for c, (t, w) in enumerate(zip(titulos, larguras), start=1):
        cel = ws.cell(row=1, column=c, value=t)
        cel.font, cel.fill = F_CABEC, PREENCHE_CABEC
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cel.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"


def montar_planilha(pares, programas, destino: Path):
    cand = [p for p in pares if p["tipo"]]
    desc = [p for p in pares if not p["tipo"]]

    wb = Workbook()

    # --- Legenda
    ws = wb.active
    ws.title = "Legenda"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    linhas = [
        ("Gabarito (ground-truth): pares candidatos de comentário e código", None),
        ("", None),
        ("Origem", "Programas da CBSA na pasta ORIGINAL do repositório da base de dados (commit e22475691524ef72a92ee77f7040233e098ab1b9, ijmitch/cics-banking-sample-application-cbsa). Números de linha referem-se aos arquivos do ORIGINAL."),
        ("Como foi gerado", "Extração automática (script extrair_pares_gabarito.py). Cada bloco de comentário da PROCEDURE DIVISION é pareado ao código que o segue, até o próximo comentário (limite de %d linhas por trecho)." % LIMITE_LINHAS),
        ("Filtro (E4, seção 1.3)", "Retido: trecho com IF, EVALUATE ou WHEN, ou com tratamento de falha (ABEND, SQLCODE, HANDLE, DFHRESP, EIBRESP, EIBRCODE). Os demais ficam na aba Descartados, para que o critério seja auditável."),
        ("O que preencher", "Aba Candidatos, colunas J a M (fundo amarelo): Revisão discente, Revisão orientador, Regra atômica (redação final) e Observações. Na aba Descartados, coluna I."),
        ("Valores da revisão", "Manter, Descartar ou Ajustar (lista suspensa). Em Descartados: Manter descartado ou Reincluir."),
        ("Alertas", "possível código desativado: o comentário parece uma instrução COBOL comentada. comentário curto: até 3 palavras. trecho truncado: o código do par passa do limite de linhas; conferir o restante no fonte. comentários adjacentes fundidos: o bloco seguinte de comentário foi unido porque o trecho entre eles não tinha código ativo. comentário menciona condição (aba Descartados): o comentário descreve uma condição, mas o trecho seguinte é só ação; a decisão pode estar antes no fonte."),
        ("Limitações", "A triagem é heurística e não substitui a revisão dupla. Comentários de cabeçalho (antes da PROCEDURE DIVISION) não entram. Linhas de depuração (indicador D) aparecem com o prefixo [D]. O tipo (Decisão, Falha, Decisão e falha) é atribuído ao trecho inteiro, não ao comentário isolado."),
        ("", None),
        ("Exemplo de regra atômica", "Ilustrativo, não faz parte do gabarito: \"Se o valor solicitado for um débito e o saldo disponível for insuficiente, a solicitação é negada.\" (uma condição e uma consequência por regra)"),
    ]
    for r, (a, b) in enumerate(linhas, start=1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = F_TITULO if r == 1 else F_NEGRITO
        ca.alignment = Alignment(vertical="top", wrap_text=(r != 1))
        if b is not None:
            cb = ws.cell(row=r, column=2, value=b)
            cb.font, cb.alignment = F_NORMAL, TOPO

    # --- Resumo (fórmulas)
    rs = wb.create_sheet("Resumo")
    titulos = [
        "Programa", "Blocos com texto", "Candidatos retidos", "Descartados",
        "Decisão", "Falha", "Decisão e falha", "Com alertas",
        "Manter (discente)", "Descartar (discente)", "Ajustar (discente)", "Pendentes (discente)",
    ]
    cabecalho(rs, titulos, [14, 12, 12, 12, 10, 10, 12, 11, 12, 12, 12, 12])
    rs.freeze_panes = "B2"
    for r, prog in enumerate(programas, start=2):
        rs.cell(row=r, column=1, value=prog)
        rs.cell(row=r, column=2, value=f"=C{r}+D{r}")
        rs.cell(row=r, column=3, value=f"=COUNTIF(Candidatos!$B:$B,$A{r})")
        rs.cell(row=r, column=4, value=f"=COUNTIF(Descartados!$B:$B,$A{r})")
        for col, tipo in ((5, "Decisão"), (6, "Falha"), (7, "Decisão e falha")):
            rs.cell(row=r, column=col, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$G:$G,"{tipo}")')
        rs.cell(row=r, column=8, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$I:$I,"<>")')
        for col, val in ((9, "Manter"), (10, "Descartar"), (11, "Ajustar")):
            rs.cell(row=r, column=col, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$J:$J,"{val}")')
        rs.cell(row=r, column=12, value=f"=C{r}-I{r}-J{r}-K{r}")
    total = len(programas) + 2
    rs.cell(row=total, column=1, value="TOTAL")
    for col in range(2, 13):
        letra = rs.cell(row=1, column=col).column_letter
        rs.cell(row=total, column=col, value=f"=SUM({letra}2:{letra}{total - 1})")
    for row in rs.iter_rows(min_row=2, max_row=total):
        for c in row:
            c.font = F_NEGRITO if c.row == total else F_NORMAL
    rs.cell(row=total + 2, column=1, value="Contagens por fórmula: atualizam conforme a revisão é preenchida na aba Candidatos.").font = F_NORMAL

    # --- Candidatos
    wc = wb.create_sheet("Candidatos")
    cabecalho(
        wc,
        ["ID", "Programa", "Linhas do comentário", "Linhas do código", "Comentário", "Trecho de código",
         "Tipo", "Estruturas", "Alertas", "Revisão discente", "Revisão orientador",
         "Regra atômica (redação final)", "Observações"],
        [13, 10, 11, 11, 60, 72, 15, 15, 24, 13, 13, 50, 30],
    )
    for r, p in enumerate(cand, start=2):
        vals = [p["id"], p["programa"], p["l_com"], p["l_cod"], p["comentario"], p["codigo"],
                p["tipo"], p["estruturas"], p["alertas"], None, None, None, None]
        for c, v in enumerate(vals, start=1):
            cel = wc.cell(row=r, column=c, value=v)
            cel.font = F_CODIGO if c == 6 else F_NORMAL
            cel.alignment = TOPO
            if c >= 10:
                cel.fill = PREENCHE_ENTRADA
    if cand:
        ultima = len(cand) + 1
        dv = DataValidation(type="list", formula1='"Manter,Descartar,Ajustar"', allow_blank=True)
        wc.add_data_validation(dv)
        dv.add(f"J2:K{ultima}")
        wc.auto_filter.ref = f"A1:M{ultima}"

    # --- Descartados
    wd = wb.create_sheet("Descartados")
    cabecalho(
        wd,
        ["ID", "Programa", "Linhas do comentário", "Linhas do código", "Comentário", "Trecho de código",
         "Motivo do descarte", "Alertas", "Reavaliação"],
        [13, 10, 11, 11, 60, 72, 30, 24, 16],
    )
    for r, p in enumerate(desc, start=2):
        motivo = "Sem código associado" if p["sem_codigo"] else "Procedural (sem decisão nem tratamento de falha)"
        alertas = [p["alertas"]] if p["alertas"] else []
        if RE_CONDICAO.search(p["comentario"]):
            alertas.append("comentário menciona condição: reavaliar")
        vals = [p["id"], p["programa"], p["l_com"], p["l_cod"], p["comentario"], p["codigo"],
                motivo, "; ".join(alertas) or None, None]
        for c, v in enumerate(vals, start=1):
            cel = wd.cell(row=r, column=c, value=v)
            cel.font = F_CODIGO if c == 6 else F_NORMAL
            cel.alignment = TOPO
            if c == 9:
                cel.fill = PREENCHE_ENTRADA
    if desc:
        dv2 = DataValidation(type="list", formula1='"Manter descartado,Reincluir"', allow_blank=True)
        wd.add_data_validation(dv2)
        dv2.add(f"I2:I{len(desc) + 1}")
        wd.auto_filter.ref = f"A1:I{len(desc) + 1}"

    wb.save(destino)
    return len(cand), len(desc)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--origem", required=True, type=Path, help="pasta cobol_src (somente leitura)")
    ap.add_argument("--saida", required=True, type=Path, help="arquivo .xlsx de saída")
    ap.add_argument("--programas", nargs="+", default=PROGRAMAS_PADRAO, help="programas a processar (sem extensão)")
    args = ap.parse_args()

    if not args.origem.is_dir():
        sys.exit(f"Origem não encontrada: {args.origem}")

    pares, ok = [], []
    for prog in args.programas:
        achados = [p for p in args.origem.iterdir() if p.stem.upper() == prog.upper() and p.suffix.lower() in {".cbl", ".cob"}]
        if not achados:
            sys.exit(f"Programa não encontrado na origem: {prog}")
        pares += extrair(achados[0])
        ok.append(prog.upper())

    n_cand, n_desc = montar_planilha(pares, ok, args.saida)
    print(f"{len(ok)} programas | {len(pares)} blocos com texto | {n_cand} candidatos | {n_desc} descartados -> {args.saida}")
    for prog in ok:
        c = sum(1 for p in pares if p["programa"] == prog and p["tipo"])
        d = sum(1 for p in pares if p["programa"] == prog and not p["tipo"])
        print(f"  {prog:9} candidatos={c:3}  descartados={d:3}")


if __name__ == "__main__":
    main()
