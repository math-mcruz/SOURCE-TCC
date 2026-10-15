#!/usr/bin/env python3
"""
migrar_revisao.py

Copia a revisão manual de uma planilha antiga para a planilha nova gerada por
monta_planilha.py, casando as linhas pelo ID. Não altera as duas planilhas de
entrada; grava uma terceira.

O que é copiado (por nome de coluna, então funciona com layouts diferentes):
  Candidatos:  Regra de negócio extraída, Categoria, Status, Observação,
               ID atômico (sub-regra) e Comentário defasado? (se existirem na antiga)
  Descartados: Reavaliação

Nada é decidido pelo script. Tudo o que muda de situação entra na aba
"Migração" da planilha de saída, para você conferir:
  - trecho de código ou tipo diferente do que foi revisado;
  - candidato que antes era descartado (e o que foi marcado na época);
  - descartado que antes era candidato (revisão que não tem onde ficar);
  - reincluído manualmente que agora o script já retém;
  - reincluído manualmente que o script ainda descarta: a linha reincluída é
    anexada ao fim de Candidatos, com os dados da planilha antiga;
  - IDs que existem só em uma das planilhas.

Programas que não estão na planilha nova (por exemplo, não reprocessados) são
ignorados, e as linhas deles na antiga não são tocadas nem copiadas.

Uso:
  python migrar_revisao.py --antiga gabarito_candidatos_rascunho_v2.xlsx --nova gabarito_candidatos_v3.xlsx --saida gabarito_candidatos_v3_migrado.xlsx
"""

import argparse
import sys
from pathlib import Path

import re
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUNAS_CAND = [
    "Regra de negócio extraída", "Categoria", "Status", "Observação",
    "ID atômico (sub-regra)", "Comentário defasado?",
]
COLUNAS_DESC = ["Reavaliação"]


def tabela(ws):
    linhas = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(linhas[0]) if h}
    return idx, [l for l in linhas[1:] if l and l[0]]


def por_id(idx, linhas):
    d = {}
    for l in linhas:
        d.setdefault(l[idx["ID"]], []).append(l)
    return d


def norm(v):
    return " ".join(str(v).split()) if v is not None else ""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--antiga", required=True, type=Path)
    ap.add_argument("--nova", required=True, type=Path)
    ap.add_argument("--saida", required=True, type=Path)
    ap.add_argument("--forcar", action="store_true")
    a = ap.parse_args()

    for p in (a.antiga, a.nova):
        if not p.is_file():
            sys.exit(f"Arquivo não encontrado: {p}")
    if a.saida.exists() and not a.forcar:
        sys.exit(f"A saída já existe e não foi sobrescrita: {a.saida} (use outro nome ou --forcar)")
    if a.saida.resolve() in (a.antiga.resolve(), a.nova.resolve()):
        sys.exit("A saída não pode ser igual a uma das entradas.")

    wb_a = load_workbook(a.antiga, data_only=True)
    wb_n = load_workbook(a.nova)
    ia_c, la_c = tabela(wb_a["Candidatos"])
    ia_d, la_d = tabela(wb_a["Descartados"])
    old_c, old_d = por_id(ia_c, la_c), por_id(ia_d, la_d)
    # linha (na aba antiga) de cada ID de Candidatos, para copiar a cor da fonte (azul = rascunho)
    ws_ac = wb_a["Candidatos"]
    linha_c = {}
    for rr in range(2, ws_ac.max_row + 1):
        v = ws_ac.cell(row=rr, column=ia_c["ID"] + 1).value
        if v and v not in linha_c:
            linha_c[v] = rr

    wc, wd = wb_n["Candidatos"], wb_n["Descartados"]
    in_c, ln_c = tabela(wc)
    in_d, ln_d = tabela(wd)
    programas = {l[in_c["Programa"]] for l in ln_c} | {l[in_d["Programa"]] for l in ln_d}
    ids_novos = {l[in_c["ID"]] for l in ln_c} | {l[in_d["ID"]] for l in ln_d}

    relato = []  # (ID, programa, situação, detalhe)
    copiadas = 0
    a_anexar = []  # linhas "Reincluído" da antiga cujo ID segue descartado na nova

    # --- Candidatos
    for r, l in enumerate(ln_c, start=2):
        id_ = l[in_c["ID"]]
        prog = l[in_c["Programa"]]
        if id_ in old_c:
            antiga = old_c[id_][0]
            for col in COLUNAS_CAND:
                if col in ia_c and col in in_c and antiga[ia_c[col]] not in (None, ""):
                    cel = wc.cell(row=r, column=in_c[col] + 1, value=antiga[ia_c[col]])
                    cel.font = copy(ws_ac.cell(row=linha_c[id_], column=ia_c[col] + 1).font)
            copiadas += 1
            tipo_a = antiga[ia_c["Tipo"]]
            status_a = antiga[ia_c["Status"]] if "Status" in ia_c else None
            if tipo_a == "Reincluído":
                relato.append((id_, prog, "reincluído manualmente, agora retido pelo script",
                               f"Revisão copiada. Tipo agora: {l[in_c['Tipo']]}. Confira se a linha reincluída antiga não ficou duplicada na planilha antiga."))
                continue
            motivos = []
            if norm(antiga[ia_c["Linhas do código"]]) != norm(l[in_c["Linhas do código"]]):
                motivos.append(f"linhas do código: antes {antiga[ia_c['Linhas do código']]}, agora {l[in_c['Linhas do código']]}")
            if norm(antiga[ia_c["Comentário"]]) != norm(l[in_c["Comentário"]]):
                motivos.append("texto do comentário diferente")
            if tipo_a != l[in_c["Tipo"]]:
                motivos.append(f"tipo: antes {tipo_a}, agora {l[in_c['Tipo']]}")
            if motivos:
                relato.append((id_, prog, "trecho ou tipo mudou; reconferir a revisão",
                               "; ".join(motivos) + f". Status copiado: {status_a}."))
        elif id_ in old_d:
            marca = old_d[id_][0][ia_d["Reavaliação"]] if "Reavaliação" in ia_d else None
            relato.append((id_, prog, "candidato novo (antes descartado)",
                           f"Reavaliação na planilha antiga: {marca or 'em branco'}. Sem revisão a copiar."))
        else:
            relato.append((id_, prog, "ID só na planilha nova", "Sem revisão a copiar."))

    # --- Descartados
    for r, l in enumerate(ln_d, start=2):
        id_ = l[in_d["ID"]]
        prog = l[in_d["Programa"]]
        if id_ in old_d and "Reavaliação" in ia_d and "Reavaliação" in in_d:
            v = old_d[id_][0][ia_d["Reavaliação"]]
            if v not in (None, ""):
                wd.cell(row=r, column=in_d["Reavaliação"] + 1, value=v)
        reinc = [x for x in old_c.get(id_, []) if x[ia_c["Tipo"]] == "Reincluído"]
        if reinc:
            a_anexar.append((id_, prog, reinc[0]))
            relato.append((id_, prog, "reincluído manualmente, mantido (anexado ao fim de Candidatos)",
                           "O script ainda descarta este trecho; a linha reincluída da planilha antiga foi copiada como estava."))
        elif id_ in old_c:
            antiga = old_c[id_][0]
            status_a = antiga[ia_c["Status"]] if "Status" in ia_c else None
            relato.append((id_, prog, "era candidato, agora descartado",
                           f"Revisão antiga sem lugar na nova (Categoria: {antiga[ia_c.get('Categoria', 0)] if 'Categoria' in ia_c else '-'}; Status: {status_a}). Decida se reinclui."))
        elif id_ not in old_d:
            relato.append((id_, prog, "ID só na planilha nova", "Sem revisão a copiar."))

    # --- anexa as reincluídas que o script segue descartando
    if a_anexar:
        modelo = wc.max_row
        nomes_novos = list(in_c)
        for id_, prog, antiga in a_anexar:
            r = wc.max_row + 1
            for nome, col in in_c.items():
                if nome in ia_c:
                    v = antiga[ia_c[nome]]
                    if v not in (None, ""):
                        wc.cell(row=r, column=col + 1, value=v)
            for col in range(1, len(in_c) + 1):
                wc.cell(row=r, column=col)._style = copy(wc.cell(row=modelo, column=col)._style)
            for nome in COLUNAS_CAND:
                if nome in ia_c and nome in in_c:
                    wc.cell(row=r, column=in_c[nome] + 1).font = copy(ws_ac.cell(row=linha_c[id_], column=ia_c[nome] + 1).font)
            for dv in wc.data_validations.dataValidation:
                m = re.match(r"([A-Z]+)\d+", str(dv.sqref).split()[0])
                if m:
                    dv.add(f"{m.group(1)}{r}")
        ultima_col = wc.cell(row=1, column=len(in_c)).column_letter
        wc.auto_filter.ref = f"A1:{ultima_col}{wc.max_row}"

    # --- IDs que sumiram (só dos programas processados na nova)
    for id_ in sorted(set(old_c) | set(old_d)):
        linha = (old_c.get(id_) or old_d.get(id_))[0]
        idx = ia_c if id_ in old_c else ia_d
        prog = linha[idx["Programa"]]
        if prog in programas and id_ not in ids_novos:
            relato.append((id_, prog, "ID só na planilha antiga", "Não existe na nova."))

    # --- aba Migração
    if "Migração" in wb_n.sheetnames:
        del wb_n["Migração"]
    wm = wb_n.create_sheet("Migração")
    cab = ["ID", "Programa", "Situação", "Detalhe"]
    larg = [14, 10, 48, 110]
    for c, (t, w) in enumerate(zip(cab, larg), start=1):
        cel = wm.cell(row=1, column=c, value=t)
        cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="2F4F6F")
        wm.column_dimensions[cel.column_letter].width = w
    wm.freeze_panes = "A2"
    for r, linha in enumerate(sorted(relato, key=lambda x: (x[2], x[0])), start=2):
        for c, v in enumerate(linha, start=1):
            cel = wm.cell(row=r, column=c, value=v)
            cel.font = Font(name="Arial", size=10)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
    wm.auto_filter.ref = f"A1:D{max(2, len(relato) + 1)}"

    wb_n.save(a.saida)

    print(f"Linhas de Candidatos com revisão copiada: {copiadas}")
    resumo = {}
    for _, _, sit, _ in relato:
        resumo[sit] = resumo.get(sit, 0) + 1
    print(f"Itens para conferir na aba Migração: {len(relato)}")
    for sit, n in sorted(resumo.items()):
        print(f"  {n:4}  {sit}")
    print(f"Gravado em {a.saida}")


if __name__ == "__main__":
    main()
