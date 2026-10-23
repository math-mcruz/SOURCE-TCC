#!/usr/bin/env python3
"""
monta_planilha.py (versão revisada)

Passo 2 do método (E4, seção 1.3): extração e filtragem dos pares
comentário-código da CBSA, para revisão conjunta discente-orientador.

Procedimento (heurístico, somente leitura na origem):
  1. Localiza a PROCEDURE DIVISION de cada programa (comentários de cabeçalho,
     anteriores a ela, não entram).
  2. Cada bloco contíguo de comentário (formato fixo: '*' ou '/' na coluna 7)
     é pareado ao código que o segue, até o próximo bloco de comentário com
     texto. O trecho EXIBIDO é limitado a LIMITE_LINHAS linhas; a CLASSIFICAÇÃO
     usa sempre o bloco inteiro.
  3. Retém o par se o bloco contém estrutura de decisão (IF, EVALUATE, WHEN)
     ou tratamento de falha (ABEND, SQLCODE, HANDLE, DFHRESP, EIBRESP, EIBRCODE).
     Os demais vão para a aba "Descartados" (rastreabilidade do filtro).
  4. Cabeçalho de decisão: quando o comentário vem logo após uma linha IF,
     ELSE, WHEN ou EVALUATE (o comentário explica o corpo do ramo), essa linha
     pertence ao código do par anterior. Se o corpo do par não tiver decisão
     nem falha, o cabeçalho é usado para reter o par e é exibido com o prefixo
     [cab.]. Pares já retidos não mudam.
  5. Gera a planilha de revisão (.xlsx) com colunas para revisão do discente,
     do orientador e a redação final da regra atômica.

Verificações de segurança:
  - aborta se o fonte tiver tabulação (desloca as colunas do formato fixo);
  - aborta se não achar a PROCEDURE DIVISION;
  - não sobrescreve um .xlsx existente (use --forcar para permitir);
  - imprime a cobertura (linhas de código dentro e fora de pares);
  - --comparar planilha_antiga.xlsx lista o que mudou de aba/tipo por ID;
  - registra o SHA-256 dos fontes na aba Legenda.

Uso (Windows):
  pip install openpyxl
  python monta_planilha.py --origem "...\\DATABASE\\ORIGINAL\\DATABASE\\cobol_src" --saida gabarito_candidatos_v3.xlsx
  python monta_planilha.py --origem "..." --saida novo.xlsx --comparar gabarito_candidatos_rascunho.xlsx

Opcional: --programas DBCRFUN XFRFUN BNK1CRA BNK1TFN  (padrão: esses quatro)
"""

import argparse
import hashlib
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

LIMITE_LINHAS = 50          # só para EXIBIÇÃO do trecho; a classificação usa o bloco inteiro
MAX_LINHAS_CONDICAO = 6     # máximo de linhas de uma condição multilinha ao buscar o cabeçalho
PROGRAMAS_PADRAO = ["DBCRFUN", "XFRFUN", "BNK1CRA", "BNK1TFN"]
NOME_SCRIPT = Path(sys.argv[0]).name if sys.argv and sys.argv[0] else "monta_planilha.py"

RE_PD = re.compile(r"^\s*PROCEDURE\s+DIVISION")
RE_LIT = re.compile(r"'[^']*'|\"[^\"]*\"")
RE_IF = re.compile(r"(?<![-\w])IF(?![-\w])")
RE_EVAL = re.compile(r"(?<![-\w])EVALUATE(?![-\w])")
RE_WHEN = re.compile(r"(?<![-\w])WHEN(?![-\w])")
RE_FALHA = re.compile(r"(?<![-\w])(ABEND|SQLCODE|HANDLE|DFHRESP|EIBRESP|EIBRCODE)(?![-\w])")
RE_CAB = re.compile(r"^(IF|ELSE|WHEN|EVALUATE)(?![-\w])")
RE_CONT = re.compile(r"^(?:(?:AND|OR|NOT)(?![-\w])|[<>=])")
RE_PERFORM = re.compile(r"(?<![-\w])PERFORM(?![-\w])")
RE_GOTO = re.compile(r"(?<![-\w])GO\s+TO(?![-\w])")
RE_PRINT = re.compile(r"^\s*(EJECT|SKIP1|SKIP2|SKIP3)\b", re.IGNORECASE)  # diretivas de listagem, não são código
RE_CONDICAO = re.compile(
    r"\b(if|when|unless|only|must|should|ensure|check|otherwise|cannot|not|sufficient|invalid|valid|"
    r"else|did|successful|unsuccessful|fail|failed|failure|worked|greater|less|equal|whether|anything|something)\b",
    re.IGNORECASE,
)
RE_DESATIVADO = re.compile(
    r"^(MOVE|DISPLAY|PERFORM|EXEC|SET|COMPUTE|ADD|SUBTRACT|CALL|GO\s+TO|INITIALIZE|"
    r"END-\w+|IF|EVALUATE|SOURCE-COMPUTER|OBJECT-COMPUTER)\b"
)


class ErroFonte(Exception):
    """Fonte que não permite extração confiável."""


def eh_comentario(l):
    return len(l) > 6 and l[6] in "*/"


def eh_debug(l):
    return len(l) > 6 and l[6] in "Dd"


def texto_comentario(l):
    """Texto de uma linha de comentário, sem indicador nem faixas de asteriscos."""
    return l[7:].strip().lstrip("*").strip()


def intervalo(a, b):
    return f"{a}" if a == b else f"{a}-{b}"


def contar(linhas_codigo):
    """Conta (IF, EVALUATE, WHEN, falha) nas linhas de código ativas."""
    n_if = n_ev = n_wh = 0
    falha = False
    for l in linhas_codigo:
        if eh_debug(l):
            continue
        t = RE_LIT.sub("''", l[7:72])
        n_if += len(RE_IF.findall(t))
        n_ev += len(RE_EVAL.findall(t))
        n_wh += len(RE_WHEN.findall(t))
        falha = falha or bool(RE_FALHA.search(t))
    return n_if, n_ev, n_wh, falha


def classificar(linhas_codigo):
    """Retorna (tipo, estruturas) para o trecho; tipo None = procedural."""
    n_if, n_ev, n_wh, falha = contar(linhas_codigo)
    estruturas = []
    if n_if:
        estruturas.append("IF")
    if n_ev:
        estruturas.append("EVALUATE")
    if n_wh:
        estruturas.append("WHEN")
    if falha:
        estruturas.append("falha")
    decisao = (n_if + n_ev + n_wh) > 0
    if decisao and falha:
        tipo = "Decisão e falha"
    elif decisao:
        tipo = "Decisão"
    elif falha:
        tipo = "Falha"
    else:
        tipo = None
    return tipo, ", ".join(estruturas)


def achar_cabecalho(L, i, limite_inf):
    """Índices das linhas do cabeçalho de decisão (IF, ELSE, WHEN, EVALUATE)
    imediatamente anterior ao bloco de comentário que começa em i, ou [].

    Pula linhas em branco e comentários vazios ('*'). Aceita condição em várias
    linhas (continuação iniciada por AND, OR, NOT ou operador), até
    MAX_LINHAS_CONDICAO linhas. Qualquer outra instrução no caminho cancela."""
    m = i - 1
    while m >= limite_inf and (not L[m].strip() or (eh_comentario(L[m]) and not texto_comentario(L[m]))):
        m -= 1
    linhas = []
    while len(linhas) < MAX_LINHAS_CONDICAO:
        if m < limite_inf or not L[m].strip() or eh_comentario(L[m]) or eh_debug(L[m]):
            return []
        t = L[m][7:72].strip()
        linhas.append(m)
        if RE_CAB.match(t):
            return linhas[::-1]
        if not RE_CONT.match(t):
            return []
        m -= 1
    return []


def extrair(caminho: Path):
    """Retorna (pares, info). Levanta ErroFonte se o fonte não for confiável."""
    programa = caminho.stem.upper()
    bruto = caminho.read_bytes()
    dados = bruto.decode("latin-1").replace("\r", "")
    if "\t" in dados:
        raise ErroFonte(
            f"{caminho.name}: contém tabulação; no formato fixo isso desloca as colunas e a "
            "detecção de comentários deixa de ser confiável. Converta as tabulações antes."
        )
    L = dados.split("\n")
    if L and L[-1] == "":
        L.pop()
    n = len(L)
    pd = next((i for i, l in enumerate(L) if not eh_comentario(l) and RE_PD.search(l[7:72])), None)
    if pd is None:
        raise ErroFonte(f"{caminho.name}: PROCEDURE DIVISION não encontrada.")

    pares = []
    cobertas = set()
    nao_exibidas = 0
    seq = 0
    i = pd + 1
    while i < n:
        if not eh_comentario(L[i]):
            i += 1
            continue
        j = i
        while j < n and (eh_comentario(L[j]) or not L[j].strip()):
            j += 1
        com = [(k, texto_comentario(L[k])) for k in range(i, j) if eh_comentario(L[k])]
        com = [(k, t) for k, t in com if t]
        if not com:
            i = j
            continue

        cab = achar_cabecalho(L, i, pd + 1)

        fundidos = False
        cod = []
        while True:
            k = j
            while k < n:
                if eh_comentario(L[k]):
                    if texto_comentario(L[k]):
                        break
                    k += 1
                    continue
                if L[k].strip() and not RE_PRINT.match(L[k][7:72]):
                    cod.append(k)
                k += 1
            ativo = any(not eh_debug(L[m]) for m in cod)
            if ativo or k >= n:
                break
            # sem código ativo (só depuração ou nada): o comentário seguinte
            # continua a mesma explicação; funde os blocos
            j2 = k
            while j2 < n and (eh_comentario(L[j2]) or not L[j2].strip()):
                j2 += 1
            extra = [(m, texto_comentario(L[m])) for m in range(k, j2) if eh_comentario(L[m])]
            com += [(m, t) for m, t in extra if t]
            fundidos = True
            j = j2

        # a lista completa serve à classificação; só a exibição é limitada
        truncado = len(cod) > LIMITE_LINHAS
        exibidas = cod[:LIMITE_LINHAS]
        cobertas.update(cod)
        nao_exibidas += len(cod) - len(exibidas)

        seq += 1
        texto = " ".join(t for _, t in com)
        alertas = []
        if RE_DESATIVADO.match(com[0][1]):
            alertas.append("possível código desativado")
        if len(texto.split()) <= 3:
            alertas.append("comentário curto")
        if fundidos:
            alertas.append("comentários adjacentes fundidos")
        if truncado:
            alertas.append(f"trecho truncado (exibidas {LIMITE_LINHAS} de {len(cod)} linhas)")

        if cod:
            tipo, estruturas = classificar([L[m] for m in cod])
            n_if, n_ev, n_wh, _ = contar([L[m] for m in cod])
            qtd_decisoes = n_if + n_ev + n_wh
            texto_cod = [RE_LIT.sub("''", L[m][7:72]) for m in cod if not eh_debug(L[m])]
            if tipo is None and any(RE_PERFORM.search(t) for t in texto_cod):
                alertas.append("comentário delega lógica (PERFORM)")
            if any(RE_GOTO.search(t) for t in texto_cod):
                alertas.append("contém GO TO (desvio incondicional)")
            por_cabecalho = False
            if tipo is None and cab:
                palavra = RE_CAB.match(L[cab[0]][7:72].strip()).group(1)
                tipo, estruturas = "Decisão", f"{palavra} (cabeçalho)"
                por_cabecalho = True
                qtd_decisoes += 1
                alertas.append(
                    f"decisão no cabeçalho anterior ao comentário (linha {intervalo(cab[0] + 1, cab[-1] + 1)})"
                )
            linhas_txt = []
            if por_cabecalho:
                linhas_txt += ["[cab.] " + L[m][7:72].rstrip() for m in cab]
            linhas_txt += [("[D] " if eh_debug(L[m]) else "") + L[m][7:72].rstrip() for m in exibidas]
            codigo = "\n".join(linhas_txt)
            l_cod = intervalo(cod[0] + 1, cod[-1] + 1)
            if por_cabecalho:
                l_cod = f"{intervalo(cab[0] + 1, cab[-1] + 1)}; {l_cod}"
        else:
            tipo, estruturas, codigo, l_cod, qtd_decisoes = None, "", "", "", 0

        pares.append(
            {
                "id": f"{programa}-{seq:03d}",
                "programa": programa,
                "l_com": intervalo(com[0][0] + 1, com[-1][0] + 1),
                "l_cod": l_cod,
                "comentario": texto,
                "codigo": codigo,
                "tipo": tipo,
                "estruturas": estruturas,
                "qtd_decisoes": qtd_decisoes,
                "alertas": "; ".join(alertas) or None,
                "sem_codigo": not cod,
            }
        )
        i = k

    ativas = [k for k in range(pd + 1, n)
              if L[k].strip() and not eh_comentario(L[k]) and not RE_PRINT.match(L[k][7:72])]
    fora = [k for k in ativas if k not in cobertas]
    n_if, n_ev, n_wh, _ = contar([L[k] for k in fora])
    info = {
        "programa": programa,
        "arquivo": caminho.name,
        "sha256": hashlib.sha256(bruto).hexdigest(),
        "ativas": len(ativas),
        "em_pares": len(cobertas),
        "fora": len(fora),
        "decisoes_fora": n_if + n_ev + n_wh,
        "nao_exibidas": nao_exibidas,
    }
    return pares, info


# ----------------------------------------------------------------- comparação
def comparar(pares, programas, antigo: Path):
    """Compara aba/tipo por ID com uma planilha anterior e imprime as diferenças."""
    wb = load_workbook(antigo, read_only=True, data_only=True)
    velho = {}
    for aba in ("Candidatos", "Descartados"):
        if aba not in wb.sheetnames:
            continue
        for row in wb[aba].iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[1] in programas:
                velho.setdefault(row[0], {})[aba] = row[6] if aba == "Candidatos" else None
    novo = {p["id"]: p for p in pares}
    diferencas = []
    for id_ in sorted(set(velho) | set(novo)):
        if id_ not in novo:
            diferencas.append(f"  {id_}: só na planilha antiga")
            continue
        p = novo[id_]
        aba_nova = "Candidatos" if p["tipo"] else "Descartados"
        if id_ not in velho:
            diferencas.append(f"  {id_}: só na nova ({aba_nova})")
            continue
        if aba_nova not in velho[id_]:
            antigas = "+".join(velho[id_])
            diferencas.append(f"  {id_}: {antigas} -> {aba_nova}" + (f" ({p['tipo']}, {p['estruturas']})" if p["tipo"] else ""))
        elif aba_nova == "Candidatos":
            t_velho = velho[id_]["Candidatos"]
            if t_velho != "Reincluído" and t_velho != p["tipo"]:
                diferencas.append(f"  {id_}: tipo {t_velho} -> {p['tipo']}")
    print(f"\nComparação com {antigo.name}: {len(diferencas)} diferença(s) em {len(set(velho) | set(novo))} ID(s)")
    for d in diferencas:
        print(d)


# ----------------------------------------------------------------- planilha
FONTE = "Arial"
F_NORMAL = Font(name=FONTE, size=10)
F_NEGRITO = Font(name=FONTE, size=10, bold=True)
F_CABEC = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
F_CODIGO = Font(name="Courier New", size=9)
F_TITULO = Font(name=FONTE, size=13, bold=True)
PREENCHE_CABEC = PatternFill("solid", fgColor="2F4F6F")
PREENCHE_ENTRADA = PatternFill("solid", fgColor="FFF9DB")
TOPO = Alignment(vertical="top", wrap_text=True)


def cabecalho(ws, titulos, larguras):
    for c, (t, w) in enumerate(zip(titulos, larguras), start=1):
        cel = ws.cell(row=1, column=c, value=t)
        cel.font, cel.fill = F_CABEC, PREENCHE_CABEC
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cel.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"


def montar_planilha(pares, programas, destino: Path, infos):
    cand = [p for p in pares if p["tipo"]]
    desc = [p for p in pares if not p["tipo"]]

    wb = Workbook()

    # --- Legenda
    ws = wb.active
    ws.title = "Legenda"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    hashes = "\n".join(f"{i['arquivo']}: {i['sha256']}" for i in infos)
    linhas = [
        ("Gabarito (ground-truth): pares candidatos de comentário e código", None),
        ("", None),
        ("Origem", "Programas da CBSA na pasta ORIGINAL do repositório da base de dados (commit e22475691524ef72a92ee77f7040233e098ab1b9, ijmitch/cics-banking-sample-application-cbsa). Números de linha referem-se aos arquivos do ORIGINAL."),
        ("Arquivos (SHA-256)", hashes),
        ("Como foi gerado", f"Extração automática (script {NOME_SCRIPT}). Cada bloco de comentário da PROCEDURE DIVISION é pareado ao código que o segue, até o próximo comentário. O trecho exibido é limitado a {LIMITE_LINHAS} linhas, mas o filtro avalia o bloco inteiro. Quando o comentário vem logo após uma linha IF, ELSE, WHEN ou EVALUATE, essa linha (cabeçalho) é usada para reter o par se o corpo não tiver decisão nem falha, e aparece com o prefixo [cab.]."),
        ("Filtro (E4, seção 1.3)", "Retido: trecho com IF, EVALUATE ou WHEN, ou com tratamento de falha (ABEND, SQLCODE, HANDLE, DFHRESP, EIBRESP, EIBRCODE). O filtro é sintático: trechos de infraestrutura pura (rollback, ABEND de DB2/VSAM, storm drain) também aparecem aqui, porque contêm essas palavras-chave; a exclusão semântica é feita na revisão manual (coluna Categoria), não pelo script. Os demais (sem decisão nem falha) ficam na aba Descartados, para que o critério seja auditável."),
        ("O que preencher", "Aba Candidatos, colunas J a M (fundo amarelo): Regra de negócio extraída, Categoria, Status e Observação, mesmas colunas usadas na revisão manual do DBCRFUN. Colunas O e P (amarelas): ID atômico (sub-regra), preenchido com a, b, c apenas quando o bloco contiver mais de uma regra, e Comentário defasado? (Sim ou Não), preenchido em toda linha revisada. Na aba Descartados, coluna I."),
        ("Valores da revisão", "Categoria: regra de negócio, infraestrutura (excluída) ou a discutir (lista suspensa). Status: A validar ou Validado. Comentário defasado?: Sim (o comentário não corresponde mais ao código) ou Não. Em Descartados: Manter descartado ou Reincluir."),
        ("Alertas", "possível código desativado: o comentário parece uma instrução COBOL comentada. comentário curto: até 3 palavras. trecho truncado: o código do par passa do limite de linhas exibidas; a coluna Linhas do código mostra o intervalo completo e a classificação considerou todo o bloco. comentários adjacentes fundidos: o bloco seguinte de comentário foi unido porque o trecho entre eles não tinha código ativo. decisão no cabeçalho anterior ao comentário: o par foi retido pela linha IF/ELSE/WHEN/EVALUATE que precede o comentário (linhas indicadas; também no código da linha anterior da planilha). comentário delega lógica (PERFORM): o trecho só chama outro parágrafo; a regra pode estar no parágrafo executado. contém GO TO: há desvio incondicional no trecho. Qtd decisões (coluna N): total de IF, EVALUATE e WHEN do bloco inteiro (+1 quando o par foi retido pelo cabeçalho); é uma contagem, não a profundidade de aninhamento. comentário menciona condição (aba Descartados): o comentário descreve uma condição, mas o trecho seguinte é só ação; a decisão pode estar antes no fonte."),
        ("Limitações", "A triagem é heurística e não substitui a revisão dupla. Comentários de cabeçalho (antes da PROCEDURE DIVISION) não entram. Regras sem comentário associado não geram par. Linhas de depuração (indicador D) aparecem com o prefixo [D]. O tipo (Decisão, Falha, Decisão e falha) é atribuído ao bloco inteiro, não ao comentário isolado. O cabeçalho IF/ELSE que precede um comentário continua também no código do par anterior (cada linha de código pertence a um único par, exceto o cabeçalho exibido como contexto)."),
        ("Rascunho (assistente)", "Colunas J, K e M preenchidas como rascunho em fonte azul; a Observação inicia por '[Rascunho]'. Status mantido em 'A validar'. Ao conferir uma linha no fonte, ajustar o texto, trocar a fonte para preta e atualizar o Status. Linhas com a categoria 'a discutir' ou 'infraestrutura (excluída)' sem regra em J são sugestões de classificação, não decisões."),
        ("Reincluídas", "Blocos marcados como 'Reincluir' na aba Descartados são registrados ao final da aba Candidatos, com o mesmo ID, Tipo 'Reincluído' e Estruturas 'procedural'. Na aba Resumo, deixam de ser contados como descartados e passam a contar como candidatos retidos (a soma de Decisão, Falha e Decisão e falha fica menor que Candidatos retidos pelo número de reincluídas)."),
        ("", None),
        ("Exemplo de regra atômica", "Ilustrativo, não faz parte do gabarito: \"Se o valor solicitado for um débito e o saldo disponível for insuficiente, a solicitação é negada.\" (uma condição e uma consequência por regra)"),
    ]
    for r, (a, b) in enumerate(linhas, start=1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = F_TITULO if r == 1 else F_NEGRITO
        ca.alignment = Alignment(vertical="top", wrap_text=(r != 1))
        if b is not None:
            cb = ws.cell(row=r, column=2, value=b)
            cb.font, cb.alignment = F_NORMAL, TOPO

    # --- Resumo (fórmulas)
    rs = wb.create_sheet("Resumo")
    titulos = [
        "Programa", "Blocos com texto", "Candidatos retidos", "Descartados",
        "Decisão", "Falha", "Decisão e falha", "Com alertas",
        "Regra de negócio", "Infraestrutura (excluída)", "A discutir", "Validados", "Pendentes",
        "Defasagem avaliada", "Comentário defasado (Sim)",
    ]
    cabecalho(rs, titulos, [14, 12, 12, 12, 10, 10, 12, 11, 13, 16, 11, 11, 11, 13, 14])
    rs.freeze_panes = "B2"
    letra_validados = rs.cell(row=1, column=12).column_letter
    for r, prog in enumerate(programas, start=2):
        rs.cell(row=r, column=1, value=prog)
        rs.cell(row=r, column=2, value=f"=C{r}+D{r}")
        rs.cell(row=r, column=3, value=f"=COUNTIF(Candidatos!$B:$B,$A{r})")
        rs.cell(row=r, column=4, value=f'=COUNTIFS(Descartados!$B:$B,$A{r},Descartados!$I:$I,"<>Reincluir")')
        for col, tipo in ((5, "Decisão"), (6, "Falha"), (7, "Decisão e falha")):
            rs.cell(row=r, column=col, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$G:$G,"{tipo}")')
        rs.cell(row=r, column=8, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$I:$I,"<>")')
        for col, val in ((9, "regra de negócio"), (10, "infraestrutura (excluída)"), (11, "a discutir")):
            rs.cell(row=r, column=col, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$K:$K,"{val}")')
        rs.cell(row=r, column=12, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$L:$L,"Validado")')
        rs.cell(row=r, column=13, value=f"=C{r}-{letra_validados}{r}")
        rs.cell(row=r, column=14, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$P:$P,"Sim")+COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$P:$P,"Não")')
        rs.cell(row=r, column=15, value=f'=COUNTIFS(Candidatos!$B:$B,$A{r},Candidatos!$P:$P,"Sim")')
    total = len(programas) + 2
    rs.cell(row=total, column=1, value="TOTAL")
    for col in range(2, 16):
        letra = rs.cell(row=1, column=col).column_letter
        rs.cell(row=total, column=col, value=f"=SUM({letra}2:{letra}{total - 1})")
    for row in rs.iter_rows(min_row=2, max_row=total):
        for c in row:
            c.font = F_NEGRITO if c.row == total else F_NORMAL
    rs.cell(row=total + 2, column=1, value="Contagens por fórmula: atualizam conforme a revisão é preenchida na aba Candidatos.").font = F_NORMAL

    # --- Candidatos
    wc = wb.create_sheet("Candidatos")
    cabecalho(
        wc,
        ["ID", "Programa", "Linhas do comentário", "Linhas do código", "Comentário", "Trecho de código",
         "Tipo", "Estruturas", "Alertas", "Regra de negócio extraída", "Categoria",
         "Status", "Observação", "Qtd decisões", "ID atômico (sub-regra)", "Comentário defasado?"],
        [13, 10, 11, 11, 60, 72, 15, 15, 24, 50, 20, 12, 30, 11, 14, 14],
    )
    for r, p in enumerate(cand, start=2):
        vals = [p["id"], p["programa"], p["l_com"], p["l_cod"], p["comentario"], p["codigo"],
                p["tipo"], p["estruturas"], p["alertas"], None, None, "A validar", None, p["qtd_decisoes"], None, None]
        for c, v in enumerate(vals, start=1):
            cel = wc.cell(row=r, column=c, value=v)
            cel.font = F_CODIGO if c == 6 else F_NORMAL
            cel.alignment = TOPO
            if 10 <= c <= 13 or c >= 15:
                cel.fill = PREENCHE_ENTRADA
    if cand:
        ultima = len(cand) + 1
        dv_cat = DataValidation(
            type="list",
            formula1='"regra de negócio,infraestrutura (excluída),a discutir"',
            allow_blank=True,
        )
        wc.add_data_validation(dv_cat)
        dv_cat.add(f"K2:K{ultima}")
        dv_status = DataValidation(type="list", formula1='"A validar,Validado"', allow_blank=True)
        wc.add_data_validation(dv_status)
        dv_status.add(f"L2:L{ultima}")
        dv_defasado = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
        wc.add_data_validation(dv_defasado)
        dv_defasado.add(f"P2:P{ultima}")
        wc.auto_filter.ref = f"A1:P{ultima}"

    # --- Descartados
    wd = wb.create_sheet("Descartados")
    cabecalho(
        wd,
        ["ID", "Programa", "Linhas do comentário", "Linhas do código", "Comentário", "Trecho de código",
         "Motivo do descarte", "Alertas", "Reavaliação"],
        [13, 10, 11, 11, 60, 72, 30, 24, 16],
    )
    for r, p in enumerate(desc, start=2):
        motivo = "Sem código associado" if p["sem_codigo"] else "Procedural (sem decisão nem tratamento de falha)"
        alertas = [p["alertas"]] if p["alertas"] else []
        if RE_CONDICAO.search(p["comentario"]):
            alertas.append("comentário menciona condição: reavaliar")
        vals = [p["id"], p["programa"], p["l_com"], p["l_cod"], p["comentario"], p["codigo"],
                motivo, "; ".join(alertas) or None, None]
        for c, v in enumerate(vals, start=1):
            cel = wd.cell(row=r, column=c, value=v)
            cel.font = F_CODIGO if c == 6 else F_NORMAL
            cel.alignment = TOPO
            if c == 9:
                cel.fill = PREENCHE_ENTRADA
    if desc:
        dv2 = DataValidation(type="list", formula1='"Manter descartado,Reincluir"', allow_blank=True)
        wd.add_data_validation(dv2)
        dv2.add(f"I2:I{len(desc) + 1}")
        wd.auto_filter.ref = f"A1:I{len(desc) + 1}"

    wb.save(destino)
    return len(cand), len(desc)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--origem", required=True, type=Path, help="pasta cobol_src (somente leitura)")
    ap.add_argument("--saida", required=True, type=Path, help="arquivo .xlsx de saída")
    ap.add_argument("--programas", nargs="+", default=PROGRAMAS_PADRAO, help="programas a processar (sem extensão)")
    ap.add_argument("--comparar", type=Path, default=None, help="planilha anterior, para listar o que mudou por ID")
    ap.add_argument("--forcar", action="store_true", help="permite sobrescrever --saida se já existir")
    args = ap.parse_args()

    if not args.origem.is_dir():
        sys.exit(f"Origem não encontrada: {args.origem}")
    if args.saida.exists() and not args.forcar:
        sys.exit(
            f"A saída já existe e não foi sobrescrita: {args.saida}\n"
            "Use outro nome em --saida (recomendado, para não perder a revisão já feita) ou --forcar."
        )
    if args.comparar is not None and not args.comparar.is_file():
        sys.exit(f"Planilha para comparação não encontrada: {args.comparar}")

    pares, ok, infos = [], [], []
    for prog in args.programas:
        achados = [p for p in args.origem.iterdir() if p.stem.upper() == prog.upper() and p.suffix.lower() in {".cbl", ".cob"}]
        if not achados:
            sys.exit(f"Programa não encontrado na origem: {prog}")
        try:
            ps, info = extrair(achados[0])
        except ErroFonte as e:
            sys.exit(f"ERRO: {e}")
        pares += ps
        infos.append(info)
        ok.append(prog.upper())

    n_cand, n_desc = montar_planilha(pares, ok, args.saida, infos)
    print(f"{len(ok)} programas | {len(pares)} blocos com texto | {n_cand} candidatos | {n_desc} descartados -> {args.saida}")
    for prog in ok:
        c = sum(1 for p in pares if p["programa"] == prog and p["tipo"])
        d = sum(1 for p in pares if p["programa"] == prog and not p["tipo"])
        print(f"  {prog:9} candidatos={c:3}  descartados={d:3}")

    print("\nCobertura (linhas de código ativas da PROCEDURE DIVISION):")
    for i in infos:
        print(
            f"  {i['programa']:9} ativas={i['ativas']:5}  em pares={i['em_pares']:5}  "
            f"fora de pares={i['fora']:4} (decisões nelas: {i['decisoes_fora']})  "
            f"não exibidas por truncamento={i['nao_exibidas']}"
        )
    print("  'fora de pares' = código antes do primeiro comentário da PROCEDURE DIVISION.")

    if args.comparar is not None:
        comparar(pares, set(ok), args.comparar)


if __name__ == "__main__":
    main()
